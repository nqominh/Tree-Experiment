"""
split_utils.py — Schema / Data Splitting

Provides split_schema_row() to separate an openpyxl sheet into a
schema (column-header) sub-sheet and a data sub-sheet.
"""

import openpyxl

from utils.sheet_utils import get_merge_cell_size, get_sub_sheet


def _is_title_row(sheet, row, ncols):
    """Check if a row is a title row (all cells have the same value or first cell spans conceptually)."""
    first_val = sheet.cell(row=row, column=1).value
    if first_val is None or str(first_val).strip() == '':
        return False

    same_value_count = 0
    for col in range(1, ncols + 1):
        cell_val = sheet.cell(row=row, column=col).value
        if cell_val == first_val:
            same_value_count += 1

    # If 80%+ of cells have the same value, it's a title row
    if same_value_count >= ncols * 0.8:
        return True
    return False


def _is_row_group_header(sheet, row, ncols):
    """Check if a row is a row group header (text in first column, empty/None in others)."""
    first_val = sheet.cell(row=row, column=1).value
    if first_val is None or str(first_val).strip() == '':
        return False

    empty_count = 0
    for col in range(2, ncols + 1):
        cell_val = sheet.cell(row=row, column=col).value
        if cell_val is None or str(cell_val).strip() in ('', 'None'):
            empty_count += 1

    data_cols = ncols - 1
    if data_cols > 0 and empty_count >= data_cols * 0.8:
        return True
    return False


def _is_data_row(sheet, row, ncols):
    """Check if a row contains data (numeric values in data columns)."""
    numeric_count = 0
    non_empty_count = 0

    for col in range(2, ncols + 1):
        cell_val = sheet.cell(row=row, column=col).value
        if cell_val is not None and str(cell_val).strip() not in ('', 'None'):
            non_empty_count += 1
            try:
                val_str = str(cell_val).replace(',', '').replace('%', '').strip()
                float(val_str)
                numeric_count += 1
            except (ValueError, TypeError):
                pass

    if numeric_count >= 3 and non_empty_count > 0:
        return True
    return False


def _is_column_header_row(sheet, row, ncols):
    """Check if a row is a column header row (distinct text values across multiple columns)."""
    if _is_row_group_header(sheet, row, ncols):
        return False

    distinct_values = set()
    non_empty_count = 0

    for col in range(1, ncols + 1):
        cell_val = sheet.cell(row=row, column=col).value
        if cell_val is not None and str(cell_val).strip() not in ('', 'None'):
            non_empty_count += 1
            distinct_values.add(str(cell_val)[:50])

    return non_empty_count >= 2 and len(distinct_values) >= 2


def _find_schema_height_by_content(sheet, nrows, ncols):
    """Find schema height by detecting where header rows end and data rows begin."""
    start_row = 1
    while start_row <= min(5, nrows) and _is_title_row(sheet, start_row, ncols):
        start_row += 1

    schema_end = start_row - 1

    for row in range(start_row, min(nrows + 1, 15)):
        if _is_data_row(sheet, row, ncols):
            return max(schema_end, row - 1)
        if _is_row_group_header(sheet, row, ncols):
            return max(schema_end, row - 1)
        if _is_column_header_row(sheet, row, ncols):
            schema_end = row

    return max(schema_end, min(4, nrows))


def split_schema_row(sheet):
    """Split sheet into (schema_sub_sheet, data_sub_sheet).

    Handles both merged cells and unmerged (flattened) sheets by using
    content-based heuristics when no merge information is present.
    """
    nrows = sheet.max_row
    ncols = sheet.max_column

    if nrows == 0 or ncols == 0:
        return (None, None)

    # Merge-based detection
    height = 1
    has_merges = False
    for col in range(1, ncols + 1):
        cell = sheet.cell(row=1, column=col)
        x1, y1, x2, y2 = get_merge_cell_size(sheet, cell.coordinate)
        if x2 > x1 or y2 > y1:
            has_merges = True
        height = max(height, x2)

    # Fallback to content-based detection for flattened sheets
    if not has_merges or height <= 1:
        height = _find_schema_height_by_content(sheet, nrows, ncols)

    height = min(height, nrows)

    return (
        get_sub_sheet(sheet, 1, 1, height, ncols),
        get_sub_sheet(sheet, height + 1, 1, nrows, ncols) if height < nrows else None,
    )
