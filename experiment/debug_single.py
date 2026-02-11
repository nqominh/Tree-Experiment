"""
debug_single.py — Interactive Single-Table Tree Debugger

Run one HTML table through all three tree-building strategies step by step,
inspecting intermediate results at each stage.

Usage:
    python -m experiment.debug_single                                # built-in sample
    python -m experiment.debug_single --file path/to/table.html
    python -m experiment.debug_single --table-id economy-table01 --data-dir RealHiTBench/html
"""

import argparse
import json
import os
import sys
import tempfile
import traceback
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from utils.sheet_utils import html2workbook, sheet2structure
from table2tree.extract_excel import get_structured_xlsx_sheet

from experiment.tree_builder import (
    html_to_tree,
    html_to_tree_direct,
    html_to_tree_fixed,
    html_to_tree_structured,
)
from experiment.tree_output import (
    tree_to_hierarchical_string,
    tree_to_schema,
    tree_to_json,
)


# ---------------------------------------------------------------------------
# Built-in sample table
# ---------------------------------------------------------------------------

SAMPLE_HTML = """\
<table>
  <tr>
    <td rowspan="2">Year</td>
    <td colspan="3">Sales</td>
    <td rowspan="2">Region</td>
  </tr>
  <tr>
    <td>Q1</td><td>Q2</td><td>Q3</td>
  </tr>
  <tr><td>2020</td><td>100</td><td>150</td><td>200</td><td>North</td></tr>
  <tr><td>2021</td><td>120</td><td>170</td><td>220</td><td>South</td></tr>
  <tr><td>2022</td><td>140</td><td>190</td><td>240</td><td>East</td></tr>
</table>
"""


# ---------------------------------------------------------------------------
# Debug step functions
# ---------------------------------------------------------------------------

def _header(msg: str):
    print(f"\n{'=' * 60}")
    print(f"  {msg}")
    print(f"{'=' * 60}")


def debug_step_html_to_excel(html: str):
    """Step 1: HTML → Excel workbook with merged cells."""
    _header("Step 1: HTML → Excel Workbook")
    wb = html2workbook(html)
    ws = wb.active
    print(f"  Rows: {ws.max_row}, Cols: {ws.max_column}")
    print(f"  Merged ranges: {[str(m) for m in ws.merged_cells.ranges]}")

    for r in range(1, ws.max_row + 1):
        vals = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            vals.append(str(v) if v is not None else ".")
        print(f"  Row {r}: {vals}")

    return wb


def debug_step_expand_merges(wb):
    """Step 2: Expand merged cells (structured sheet)."""
    _header("Step 2: Expand Merged Cells → Structured Sheet")
    tmp = tempfile.mktemp(suffix=".xlsx")
    wb.save(tmp)
    sheet = get_structured_xlsx_sheet(tmp)
    print(f"  Rows: {sheet.max_row}, Cols: {sheet.max_column}")

    for r in range(1, sheet.max_row + 1):
        vals = []
        for c in range(1, sheet.max_column + 1):
            v = sheet.cell(row=r, column=c).value
            vals.append(str(v) if v is not None else ".")
        print(f"  Row {r}: {vals}")

    return sheet


def debug_step_strategy(name: str, fn, html: str, **kwargs):
    """Try one tree-building strategy and report results."""
    _header(f"Strategy: {name}")
    try:
        tree = fn(html, **kwargs) if kwargs else fn(html)
        if tree is None:
            print("  Result: None")
            return None
        print(f"  ✅ Success — cols={tree.get_max_col()}, rows={tree.get_max_row()}")
        print(f"\n  Schema: {tree_to_schema(tree)}")
        hier = tree_to_hierarchical_string(tree)
        print(f"\n  Hierarchical string:\n{hier}")
        try:
            j = tree_to_json(tree)
            print(f"\n  JSON:\n{json.dumps(j, indent=2, ensure_ascii=False)[:1500]}")
        except Exception as e:
            print(f"  JSON failed: {e}")
        return tree
    except Exception as e:
        print(f"  ❌ Failed: {e}")
        traceback.print_exc()
        return None


def debug_table(html: str, max_header_rows: int = 2):
    """Run full debug pipeline on one HTML table."""
    _header("INPUT HTML (first 500 chars)")
    print(html[:500])

    wb = debug_step_html_to_excel(html)
    debug_step_expand_merges(wb)

    debug_step_strategy("direct", html_to_tree_direct, html)
    debug_step_strategy(
        f"fixed (max_header_rows={max_header_rows})",
        html_to_tree_fixed, html, max_header_rows=max_header_rows,
    )
    debug_step_strategy("structured", html_to_tree_structured, html)

    _header("Combined html_to_tree() [auto-fallback]")
    tree, strategy = html_to_tree(html, max_header_rows=max_header_rows)
    if tree:
        print(f"  Winner: {strategy}")
        print(f"  Cols: {tree.get_max_col()}, Rows: {tree.get_max_row()}")
    else:
        print("  All strategies failed.")

    return tree, strategy


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Debug HO-Tree building for a single table")
    parser.add_argument("--file", type=str, help="Path to an HTML file")
    parser.add_argument("--table-id", type=str, help="Table ID to look up in --data-dir")
    parser.add_argument("--data-dir", type=str, help="Directory containing HTML files")
    parser.add_argument("--header-rows", type=int, default=2, help="Max header rows for fixed strategy")
    args = parser.parse_args()

    html = None

    if args.file:
        with open(args.file, "r", encoding="utf-8") as f:
            html = f.read()
        print(f"Loaded {args.file} ({len(html)} chars)")

    elif args.table_id and args.data_dir:
        for dirpath, _, filenames in os.walk(args.data_dir):
            for fname in filenames:
                if fname.replace(".html", "") == args.table_id:
                    fpath = os.path.join(dirpath, fname)
                    with open(fpath, "r", encoding="utf-8") as f:
                        html = f.read()
                    print(f"Found {fpath} ({len(html)} chars)")
                    break
            if html:
                break
        if not html:
            print(f"ERROR: Table '{args.table_id}' not found in {args.data_dir}")
            sys.exit(1)
    else:
        print("No --file or --table-id given, using built-in sample table.\n")
        html = SAMPLE_HTML

    debug_table(html, max_header_rows=args.header_rows)


if __name__ == "__main__":
    main()
