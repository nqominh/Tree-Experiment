"""
sheet_utils.py — Excel Sheet Manipulation Utilities

Core helpers for working with openpyxl sheets: HTML conversion,
merge-cell handling, sub-sheet extraction, and structure detection.
"""

import numpy as np
from collections import Counter

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from bs4 import BeautifulSoup

from utils.constants import *


def sheet2structure(sheet):
    """Expand all merged cells so every cell contains its resolved value."""
    nrows = sheet.max_row
    ncols = sheet.max_column

    for x in range(1, nrows + 1):
        for y in range(1, ncols + 1):
            cell = sheet.cell(row=x, column=y)
            value = cell.value
            x1, y1, x2, y2 = get_merge_cell_size(sheet, cell.coordinate)
            if x1 != x2 or y1 != y2:
                sheet.unmerge_cells(start_row=x1, start_column=y1, end_row=x2, end_column=y2)

            for xx in range(x1, x2 + 1):
                for yy in range(y1, y2 + 1):
                    sheet.cell(row=xx, column=yy, value=value)
    return sheet


def html2workbook(html_content):
    """
    Convert HTML table to Excel workbook with proper rowspan/colspan handling.

    Uses an occupation grid algorithm to correctly map HTML cell positions
    to Excel coordinates when rowspan/colspan are present.
    Preserves leading whitespace/indentation (e.g., from &nbsp;) for hierarchy detection.
    """
    soup = BeautifulSoup(html_content, 'html.parser')
    table = soup.find('table')

    wb = openpyxl.Workbook()
    ws = wb.active

    # Track which cells came from <th> tags (semantic header annotation)
    th_map = {}
    ws._th_map = th_map

    # Build occupation grid for correct cell mapping with spans
    occupied = set()
    rows = table.find_all('tr')

    excel_row = 1
    for row in rows:
        excel_col = 1
        for cell in row.find_all(['td', 'th']):
            # Skip columns already occupied by a previous rowspan/colspan
            while (excel_row, excel_col) in occupied:
                excel_col += 1

            is_header = (cell.name == 'th')

            # Get text preserving leading whitespace (for indentation detection)
            cell_text = cell.get_text()
            cell_text = cell_text.replace('\xa0', ' ')  # &nbsp; → space
            cell_value = cell_text.rstrip()

            colspan = int(cell.get('colspan', 1))
            rowspan = int(cell.get('rowspan', 1))

            ws.cell(row=excel_row, column=excel_col, value=cell_value)

            if rowspan > 1 or colspan > 1:
                ws.merge_cells(
                    start_row=excel_row, start_column=excel_col,
                    end_row=excel_row + rowspan - 1, end_column=excel_col + colspan - 1
                )

            for dr in range(rowspan):
                for dc in range(colspan):
                    occupied.add((excel_row + dr, excel_col + dc))
                    th_map[(excel_row + dr, excel_col + dc)] = is_header

            excel_col += colspan
        excel_row += 1

    return wb


def delete_dict_none_none(data: dict):
    """Remove {None: None} entries from a nested dict."""
    new_dict = {}
    for k, v in data.items():
        if (k is None or k == 'None') and (v is None or v == 'None'):
            pass
        else:
            if isinstance(v, dict):
                new_dict[k] = delete_dict_none_none(v)
            elif isinstance(v, list):
                new_v = []
                for item in v:
                    if isinstance(item, dict):
                        new_v.append(delete_dict_none_none(item))
                    else:
                        new_v.append(item)
                new_dict[k] = new_v
            else:
                new_dict[k] = v
    return new_dict


def get_sub_sheet(sheet, min_row, min_col, max_row, max_col, use_wb=False):
    """Extract a rectangular sub-region of a sheet into a new workbook sheet."""
    if min_row > max_row or min_col > max_col:
        return None
    if min_row == max_row and min_col == max_col and (
        sheet.cell(row=min_row, column=min_col) is None
        or str(sheet.cell(row=min_row, column=min_col)).strip() == ""
    ):
        return None

    wb = openpyxl.Workbook()
    new_sheet = wb.active

    for i, row in enumerate(
        sheet.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
        ),
        start=1,
    ):
        for j, cell in enumerate(row, start=1):
            new_sheet.cell(row=i, column=j, value=cell.value)

    # Copy merged cell ranges that fall within the sub-region
    for merged_range in sheet.merged_cells.ranges:
        mmin_col, mmin_row, mmax_col, mmax_row = merged_range.bounds
        if (
            mmin_row >= min_row
            and mmax_row <= max_row
            and mmin_col >= min_col
            and mmax_col <= max_col
        ):
            new_sheet.merge_cells(
                get_coordinate_by_cell_pos(
                    mmin_row - min_row + 1,
                    mmin_col - min_col + 1,
                    mmax_row - min_row + 1,
                    mmax_col - min_col + 1,
                )
            )

    if use_wb:
        return wb
    return wb.active


def get_merge_cell_size(sheet, cell_pos):
    """Return the bounding box [x1, y1, x2, y2] for the merged region containing cell_pos."""
    cell = sheet[cell_pos]
    for merged_range in sheet.merged_cells.ranges:
        if cell.coordinate in merged_range:
            min_col, min_row, max_col, max_row = merged_range.bounds
            return min_row, min_col, max_row, max_col

    return get_cell_pos_by_coordinate(cell_pos)


def get_merge_cell_value(sheet, cell_pos):
    """Return the value of the top-left cell in the merged region containing cell_pos."""
    x1, y1, x2, y2 = get_merge_cell_size(sheet, cell_pos)
    return sheet.cell(row=x1, column=y1).value


def get_cell_pos_by_coordinate(coordinate_str):
    """Parse Excel coordinate string (e.g. 'A21' or 'A21:B32') into [row1, col1, row2, col2]."""
    def excel_column_to_number(column_string):
        column_number = 0
        for char in column_string:
            column_number = column_number * 26 + (ord(char.upper()) - ord("A") + 1)
        return column_number

    if ":" not in coordinate_str:
        start_col_letter = "".join(c for c in coordinate_str if c.isalpha())
        start_row = int("".join(c for c in coordinate_str if c.isdigit()))
        start_col = excel_column_to_number(start_col_letter)
        return [start_row, start_col, start_row, start_col]

    start_cell, end_cell = coordinate_str.split(":")

    start_col_letter = "".join(c for c in start_cell if c.isalpha())
    start_row = int("".join(c for c in start_cell if c.isdigit()))
    start_col = excel_column_to_number(start_col_letter)

    end_col_letter = "".join(c for c in end_cell if c.isalpha())
    end_row = int("".join(c for c in end_cell if c.isdigit()))
    end_col = excel_column_to_number(end_col_letter)

    return [start_row, start_col, end_row, end_col]


def get_coordinate_by_cell_pos(x1, y1, x2, y2):
    """Convert [row1, col1, row2, col2] into Excel coordinate string (e.g. 'A1:B2')."""
    if x1 == x2 and y1 == y2:
        return f"{get_column_letter(y1)}{x1}"
    col1 = get_column_letter(y1)
    col2 = get_column_letter(y2)
    return f"{col1}{x1}:{col2}{x2}"


def single_cell(sheet, x1, y1, x2, y2):
    """Check if the region is a single cell or a single merged cell."""
    if x1 == x2 and y1 == y2:
        return True

    for merged_range in sheet.merged_cells.ranges:
        min_c, min_r, max_c, max_r = range_boundaries(str(merged_range))
        if min_r <= x1 and min_c <= y1 and max_r >= x2 and max_c >= y2:
            return True
    return False
