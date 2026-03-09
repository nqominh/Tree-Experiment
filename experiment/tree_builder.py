"""
tree_builder.py — HO-Tree Construction from HTML

Converts HTML tables to hierarchical FeatureTree structures using
three strategies with automatic fallback:

  1. direct     — construct_sheet() with auto header detection
  2. fixed      — manual header/data split (configurable max_header_rows)
  3. structured — flat structured fallback (always succeeds)

Usage:
    from experiment.tree_builder import html_to_tree

    tree, strategy = html_to_tree(html_string)
"""

import os
import tempfile
import openpyxl

from utils.sheet_utils import html2workbook, sheet2structure
from utils.constants import DEFAULT_TABLE_NAME
from table2tree.feature_tree import (
    FeatureTree,
    IndexTree,
    BodyTree,
    IndexNode,
    BodyNode,
    construct_index_tree,
    construct_body_tree,
    construct_sheet,
    construct_feature_tree,
)
from table2tree.extract_excel import get_structured_xlsx_sheet


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _html_to_workbook_sheet(html_content: str):
    """HTML string → expanded openpyxl sheet (merged cells resolved).

    Returns (sheet, temp_path). Caller MUST call os.unlink(temp_path)
    when done to avoid leaking temp files.
    """
    wb = html2workbook(html_content)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        temp_path = tmp.name
    wb.save(temp_path)
    sheet = get_structured_xlsx_sheet(temp_path)
    return sheet, temp_path


# ---------------------------------------------------------------------------
# Strategy 1: Direct (auto header detection via split_schema_row)
# ---------------------------------------------------------------------------

def html_to_tree_direct(html_content: str) -> FeatureTree:
    """
    HTML → Excel → expand merges → construct_sheet().
    Uses ST-Raptor's split_schema_row() to auto-detect header rows.
    """
    wb = html2workbook(html_content)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        temp_path = tmp.name
    try:
        wb.save(temp_path)
        wb2 = openpyxl.load_workbook(temp_path, data_only=True)
        sheet = sheet2structure(wb2.active)
        return construct_sheet(sheet)
    finally:
        os.unlink(temp_path)


# ---------------------------------------------------------------------------
# Strategy 2: Fixed (manual header row count)
# ---------------------------------------------------------------------------

def html_to_tree_fixed(html_content: str, max_header_rows: int = 2) -> FeatureTree:
    """
    Manually split first `max_header_rows` rows as schema, remainder as data.
    More robust when auto-detection fails.
    """
    sheet, temp_path = _html_to_workbook_sheet(html_content)
    try:
        nrows = sheet.max_row
        ncols = sheet.max_column

        # Schema sheet
        schema_wb = openpyxl.Workbook()
        schema_sheet = schema_wb.active
        for r in range(1, min(max_header_rows + 1, nrows + 1)):
            for c in range(1, ncols + 1):
                schema_sheet.cell(row=r, column=c).value = sheet.cell(row=r, column=c).value

        # Data sheet
        data_wb = openpyxl.Workbook()
        data_sheet = data_wb.active
        for r in range(max_header_rows + 1, nrows + 1):
            for c in range(1, ncols + 1):
                data_sheet.cell(row=r - max_header_rows, column=c).value = sheet.cell(
                    row=r, column=c
                ).value

        index_tree = construct_index_tree(schema_sheet)
        body_tree, _ = construct_body_tree(index_tree, data_sheet)
        return FeatureTree(index_tree=index_tree, body_tree=body_tree)
    finally:
        os.unlink(temp_path)


# ---------------------------------------------------------------------------
# Strategy 3: Structured (flat, always succeeds)
# ---------------------------------------------------------------------------

def html_to_tree_structured(html_content: str) -> FeatureTree:
    """Treat full expanded sheet as structured data — last-resort fallback."""
    sheet, temp_path = _html_to_workbook_sheet(html_content)
    try:
        tree_dict = {DEFAULT_TABLE_NAME: sheet}
        return construct_feature_tree(tree_dict)
    finally:
        os.unlink(temp_path)


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def html_to_tree(html_content: str, max_header_rows: int = 2):
    """
    Build HO-Tree from an HTML table string.

    Tries strategies in order: direct → fixed → structured.

    Returns:
        (FeatureTree, strategy_name)  on success
        (None,        'failed')       if all strategies fail
    """
    strategies = [
        ("direct", lambda: html_to_tree_direct(html_content)),
        ("fixed", lambda: html_to_tree_fixed(html_content, max_header_rows)),
        ("structured", lambda: html_to_tree_structured(html_content)),
    ]
    for name, fn in strategies:
        try:
            tree = fn()
            if tree is not None and tree.get_max_col() > 0:
                return tree, name
        except Exception:
            pass
    return None, "failed"

